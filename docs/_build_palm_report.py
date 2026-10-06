from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "code" / "palm_demo" / "runtime" / "debug_capture"
ASSETS = ROOT / "docs" / "palm_report_assets"
OUTPUT = ROOT / "docs" / "2026-09-22-fastcc-two-hand-pilot-analysis.docx"
ASSETS.mkdir(exist_ok=True)

THRESHOLD = 0.28
SAME = [
    0.1885775862, 0.3064516129, 0.3448275862, 0.3559907834, 0.2558398220,
    0.3074596774, 0.3958333333, 0.2062500000, 0.1784274194,
]
DIFFERENT = [
    0.3388671875, 0.3885416667, 0.2802419355, 0.3135080645, 0.3730468750,
    0.3974654378, 0.2792338710, 0.2880859375, 0.4040948276,
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf")
    return ImageFont.truetype(str(path), size)


def make_sheet(session: str, scores: list[float], title: str, output: Path) -> None:
    width, tile_h = 1500, 325
    canvas = Image.new("RGB", (width, 120 + 5 * tile_h), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text((36, 28), title, fill="#16253A", font=font(34, True))
    for idx in range(10):
        col, row = idx % 2, idx // 2
        x, y = 30 + col * 745, 105 + row * tile_h
        sample = DATA / session / f"sample_{idx + 1:03d}"
        raw = Image.open(sample / "raw.png").convert("RGB")
        raw.thumbnail((390, 250))
        roi = Image.open(sample / "roi_128.png").convert("RGB").resize((250, 250))
        canvas.paste(raw, (x, y + 50))
        canvas.paste(roi, (x + 420, y + 50))
        if idx == 0:
            label = "01  REFERENCE"
            color = "#16253A"
        else:
            score = scores[idx - 1]
            decision = "ACCEPT" if score <= THRESHOLD else "REJECT"
            label = f"{idx + 1:02d}  distance={score:.4f}  {decision}"
            color = "#1B6E3C" if decision == "ACCEPT" else "#A62A2A"
        draw.text((x, y + 8), label, fill=color, font=font(23, True))
        draw.text((x + 420, y + 306), "128 x 128 ROI", fill="#555555", font=font(17))
    canvas.save(output, quality=94)


def make_score_plot(output: Path) -> None:
    canvas = Image.new("RGB", (1500, 850), "white")
    draw = ImageDraw.Draw(canvas)
    left, top, right, bottom = 180, 125, 1420, 700
    low, high = 0.15, 0.43
    to_y = lambda value: bottom - int((value - low) / (high - low) * (bottom - top))
    draw.text((280, 35), "Same-hand and different-hand score overlap", fill="#16253A", font=font(38, True))
    for tick in (0.15, 0.20, 0.25, 0.30, 0.35, 0.40):
        y = to_y(tick)
        draw.line((left, y, right, y), fill="#D9E1E8", width=2)
        draw.text((75, y - 14), f"{tick:.2f}", fill="#444444", font=font(22))
    draw.line((left, top, left, bottom), fill="#333333", width=3)
    draw.line((left, bottom, right, bottom), fill="#333333", width=3)
    threshold_y = to_y(THRESHOLD)
    for x in range(left, right, 28):
        draw.line((x, threshold_y, min(x + 14, right), threshold_y), fill="#222222", width=3)
    draw.text((right - 235, threshold_y - 32), "Threshold 0.28", fill="#222222", font=font(22, True))
    centers = (560, 1080)
    colors = ("#2774AE", "#D1495B")
    labels = ("Same hand", "Bankey impostor")
    for center, values, color, label in zip(centers, (SAME, DIFFERENT), colors, labels):
        offsets = (-72, -54, -36, -18, 0, 18, 36, 54, 72)
        for offset, value in zip(offsets, values):
            y = to_y(value)
            draw.ellipse((center + offset - 11, y - 11, center + offset + 11, y + 11), fill=color, outline="white", width=2)
        mean_y = to_y(sum(values) / len(values))
        draw.line((center - 115, mean_y, center + 115, mean_y), fill="#173B56" if center == centers[0] else "#842332", width=8)
        box = draw.textbbox((0, 0), label, font=font(27, True))
        draw.text((center - (box[2] - box[0]) / 2, bottom + 35), label, fill="#222222", font=font(27, True))
    draw.text((left, 82), "Fast-CC distance    lower is more similar", fill="#555555", font=font(23))
    canvas.save(output)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, "234A6F")
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "EFF5FA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        if widths:
            for c_idx, width in enumerate(widths):
                row.cells[c_idx].width = width


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            if idx > 0 and len(str(value)) < 16:
                cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_caption(doc, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.add_run(text)


def add_picture(doc, path: Path, width=Inches(6.55)) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=width)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.05)
    section.right_margin = Cm(2.05)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft JhengHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Microsoft JhengHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.size = Pt(25)
    title.font.bold = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Microsoft JhengHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Fast-CC 掌紋 ROI 兩人先導測試   2026-09-22")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def add_body(doc, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def build() -> None:
    same_sheet = ASSETS / "same_hand_contact_sheet.png"
    bankey_sheet = ASSETS / "bankey_contact_sheet.png"
    score_plot = ASSETS / "score_overlap.png"
    make_sheet("20260916T201834Z", SAME, "Test 1   Same hand against a fixed reference", same_sheet)
    make_sheet("20260916T202013Z", DIFFERENT, "Test 2   Bankey against the same fixed reference", bankey_sheet)
    make_score_plot(score_plot)

    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Fast CC 掌紋 ROI 兩人先導測試分析報告")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Raspberry Pi 5 與 Camera NoIR v2").bold = True
    date = doc.add_paragraph("2026 年 9 月 22 日")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    doc.add_heading("結論摘要", level=1)
    add_body(doc, "目前 pipeline 已能穩定輸出可追溯的 128×128 ROI，但辨識結果尚未穩定。以固定 reference 比較 9 張同手 ROI 和 9 張 Bankey 異手 ROI，在暫定門檻 0.28 下，同手只有 4 張被接受，Bankey 則有 1 張被錯誤接受。")
    add_body(doc, "核心判斷：同手 ROI 的位置和尺度變化會明顯推高 Fast-CC distance；同時，Bankey 的錯誤接受影像具有足夠 contrast 和 sharpness，因此問題也涉及 matcher 的分離能力和 provisional threshold。只調整門檻不能同時消除錯誤拒絕與錯誤接受。", "核心判斷：")

    add_table(doc,
              ["指標", "同一隻手", "Bankey 異手"],
              [
                  ["Probe 數量", "9", "9"],
                  ["平均 distance", "0.2822", "0.3403"],
                  ["Distance 範圍", "0.1784–0.3958", "0.2792–0.4041"],
                  ["門檻 0.28 下接受", "4/9", "1/9"],
                  ["門檻 0.28 下拒絕", "5/9", "8/9"],
              ], [Cm(5.2), Cm(4.4), Cm(4.4)])

    doc.add_page_break()
    doc.add_heading("1 實驗目的與協議", level=1)
    add_body(doc, "本次先導測試分開回答兩個問題：同一隻手重複放置後是否仍能被穩定接受，以及換成 Bankey 的手後是否能被穩定拒絕。這是工程診斷，不是正式的 FAR、FRR 或安全性驗證。")
    add_body(doc, "兩個測試共用 20260916T201834Z/sample_001 作 reference。這個設計避免更換 reference 導致兩組分數不可直接比較。Fast-CC distance 越低代表越相似；0.28 是現有 demo 的 provisional threshold，沒有用本批測試資料重新調校。")
    add_table(doc,
              ["用途", "Session 與樣本", "數量", "標籤"],
              [
                  ["Reference", "201834Z/sample_001", "1", "固定參考"],
                  ["同手測試", "201834Z/sample_002–010", "9", "Same hand"],
                  ["異手測試", "202013Z/sample_002–010", "9", "Bankey impostor"],
              ], [Cm(3.0), Cm(6.2), Cm(2.0), Cm(3.1)])

    doc.add_heading("2 整體結果", level=1)
    add_table(doc,
              ["真實身分", "系統接受", "系統拒絕"],
              [
                  ["同一隻手", "4  正確接受", "5  錯誤拒絕"],
                  ["Bankey 異手", "1  錯誤接受", "8  正確拒絕"],
              ], [Cm(4.4), Cm(4.8), Cm(4.8)])
    add_body(doc, "這 18 次比較的 True Accept Rate 為 44.4%，False Reject Rate 為 55.6%，True Reject Rate 為 88.9%，False Accept Rate 為 11.1%，Balanced Accuracy 為 66.7%。這些比例只描述本次 pilot，不能外推為系統整體效能。")
    add_picture(doc, score_plot, Inches(6.25))
    add_caption(doc, "圖 1  同手與 Bankey 異手的 Fast-CC distance 分布。虛線為暫定門檻 0.28，短橫線表示各組平均值。")
    add_body(doc, "同手平均 distance 為 0.2822，異手平均為 0.3403，方向符合預期；但兩組重疊區間為 0.2792–0.3958。觀察到的排序 AUC 約為 0.716，表示有一些區分訊號，但不足以支持穩定的個別判定。")

    doc.add_page_break()
    doc.add_heading("3 測試一  同一隻手重複放置", level=1)
    add_picture(doc, same_sheet, Inches(5.7))
    add_caption(doc, "圖 2  同手測試的原始影像與 128×128 ROI。所有 probe 均與 sample_001 比較。")
    add_body(doc, "同手 distance 從 0.1784 到 0.3958，跨度達 0.2174，並有 5/9 probes 被錯誤拒絕。這個跨度大於兩組平均值之間的 0.0582 差距，代表放手位置造成的變化足以蓋過部分身分差異。")
    add_body(doc, "最嚴重案例是 sample_008。Reference 的 ROI 四邊形約為 189×189 px，中心約在 (304.9, 136.3)；sample_008 約為 260×260 px，中心約在 (256.1, 188.6)。ROI 邊長增加約 37%，中心水平移動約 49 px、垂直移動約 52 px，Fast-CC distance 同時升至 0.3958。")
    add_body(doc, "這顯示輸出尺寸固定為 128×128 並不等於掌紋結構已被對齊。透視裁切前涵蓋的掌心範圍、尺度和位置仍有明顯差異。不過 sample_010 也有一定幾何差異，distance 卻只有 0.1784，因此幾何不是唯一因素；掌面姿勢、旋轉、照明與 Fast-CC 特徵敏感度仍需一併檢查。")

    doc.add_heading("4 測試二  Bankey 異手", level=1)
    add_picture(doc, bankey_sheet, Inches(5.7))
    add_caption(doc, "圖 3  Bankey 異手測試。圖中全部 distance 均使用與測試一相同的 reference。")
    add_body(doc, "Bankey 有 8/9 probes 被正確拒絕；sample_008 的 distance 為 0.2792，低於門檻 0.28，因此被錯誤接受。sample_004 為 0.2802，sample_009 為 0.2881，兩者也非常接近門檻。")
    add_body(doc, "Bankey sample_008 的 contrast 為 46.70、sharpness 為 32.83，遠高於現有最低門檻 contrast 12 和 sharpness 2。它不是一張明顯模糊或低對比的 ROI，因此不能把這次錯誤接受單純歸因於影像品質。")
    add_body(doc, "相反地，Bankey sample_003 的 contrast 只有 13.05、sharpness 6.00，接近最低要求，而且 ROI 視覺上有明顯不均勻紋理。它仍通過 READY，說明現有 quality gate 只保證最低可處理條件，不保證足以可靠辨識。")

    doc.add_heading("5 門檻與分數分離解讀", level=1)
    add_body(doc, "兩組平均值的標準化差異 Cohen's d 約為 0.87，表示平均上存在差異；但生物辨識要求每次個別判定可靠，不能只依賴平均值。兩組分布的大量重疊才是目前主要風險。")
    add_body(doc, "在本批已觀察分數上，把門檻降到約 0.2558 可以拒絕全部 Bankey probes，但同手仍只有 4/9 被接受。提高門檻則會改善同手接受率，同時增加 Bankey 被接受的風險。因此，單純調整門檻不能修復目前問題。")

    doc.add_heading("6 工程診斷", level=1)
    add_table(doc,
              ["觀察", "證據", "優先解讀"],
              [
                  ["同手高 distance", "ROI 中心與尺度明顯改變", "先改善 ROI 幾何與正規化"],
                  ["同手高 distance", "ROI 幾何正常但影像異常", "檢查曝光、清晰度與 quality gate"],
                  ["同手高 distance", "ROI 幾何與品質均正常", "檢查 Fast-CC 與 gallery 策略"],
                  ["異手低 distance", "Bankey sample_008 品質正常", "檢查 matcher 分離能力與 threshold"],
              ], [Cm(3.6), Cm(5.0), Cm(5.8)])

    doc.add_heading("7 可以支持與不能支持的結論", level=1)
    add_body(doc, "本次資料支持：pipeline 能保存完整且可追溯的 ROI；同手重複放置仍不穩定；ROI 幾何變化是具體失敗來源；Bankey 的錯誤接受顯示 matcher 和 threshold 也需要檢查；READY 只代表達到最低品質要求，不代表足以可靠辨識。")
    add_body(doc, "本次資料不能支持：正式 FAR、FRR、EER 或準確率；Fast-CC 一定不適用；ROI 是唯一失敗原因；0.28 是最佳門檻；系統具備支付級安全性；結果可泛化到更多人、不同日期、距離、角度或光線。")

    doc.add_heading("8 下一輪實驗", level=1)
    steps = [
        "同一隻手進行 30 次真正獨立放置，每次完全移開再放回，保存原圖、ROI 和 metadata。",
        "量化 ROI 中心、尺度和旋轉角的變異，並將這些幾何量與 Fast-CC distance 對照。",
        "若 ROI 已穩定而同手 distance 仍波動，改用多張 enrollment ROI，測試 median gallery score、minimum score 或 template fusion。",
        "使用獨立 development data 選門檻，保留測試資料作最終評估。",
        "基本穩定後再加入 Bankey 30 次、左右旋轉 15°/30°、不同距離、更多參與者和跨日期測試。",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    doc.add_page_break()
    doc.add_heading("附錄 A 逐張分數", level=1)
    doc.add_heading("A1 同一隻手", level=2)
    same_rows = []
    for idx, value in enumerate(SAME, 2):
        same_rows.append([f"sample_{idx:03d}", f"{value:.4f}", "ACCEPT" if value <= THRESHOLD else "REJECT", "正確" if value <= THRESHOLD else "錯誤拒絕"])
    add_table(doc, ["Probe", "Distance", "判定", "結果"], same_rows, [Cm(3.6), Cm(3.2), Cm(3.4), Cm(4.2)])

    doc.add_heading("A2 Bankey 異手", level=2)
    diff_rows = []
    for idx, value in enumerate(DIFFERENT, 2):
        diff_rows.append([f"sample_{idx:03d}", f"{value:.4f}", "ACCEPT" if value <= THRESHOLD else "REJECT", "錯誤接受" if value <= THRESHOLD else "正確"])
    add_table(doc, ["Probe", "Distance", "判定", "結果"], diff_rows, [Cm(3.6), Cm(3.2), Cm(3.4), Cm(4.2)])

    doc.add_heading("附錄 B Capture 資料結構", level=1)
    add_body(doc, "每個 debug capture session 使用 UTC 時間命名，內含 sample_001 至 sample_010。每個 sample 保存 raw.png、roi_128.png 和 metadata.json。metadata 記錄 ROI 四邊形、contrast、sharpness、追蹤狀態、相機設定、時間與檔案雜湊。sample 編號只代表保存順序，身分標籤仍需由實驗 protocol 額外記錄。")
    add_table(doc,
              ["檔案", "內容", "用途"],
              [
                  ["raw.png", "480×360 RGB 原始畫面", "檢查放手姿勢、距離、光線與 detector"],
                  ["roi_128.png", "128×128 灰階正規化 ROI", "Fast-CC 的實際輸入"],
                  ["metadata.json", "幾何、品質、時間與相機資料", "重現與故障分類"],
              ], [Cm(3.2), Cm(4.5), Cm(6.7)])

    doc.core_properties.title = "Fast CC 掌紋 ROI 兩人先導測試分析報告"
    doc.core_properties.subject = "同手重複性與 Bankey 異手比較"
    doc.core_properties.author = "FYP Palm Recognition Project"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
